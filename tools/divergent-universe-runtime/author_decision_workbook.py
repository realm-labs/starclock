"""Author a complete typed decision workbook with openpyxl; never overwrite."""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from evolution_event_rows import append_evolution_events
from sage_acquisition_rows import append_sage_acquisitions
from sage_victory_rows import append_sage_victory_blessings
from domain_choice_rows import append_domain_choices
from battle_fragment_rows import append_battle_fragments
from curio_expiry_rows import append_curio_expiries
from curio_battle_stat_rows import append_curio_battle_stats
from curio_battle_reaction_rows import append_curio_battle_reactions
from tawot_service_rows import append_tawot_services
from curio_domain_grant_rows import append_curio_domain_grants
from tawot_victory_rows import append_tawot_victory
from equation_expansion_rows import append_equation_expansion_rewards
from domain_layout_rows import append_domain_layout
from domain_deck_rows import append_domain_decks


def rows() -> dict[str, list[list[object]]]:
    revision = "fd978d6ef09f941fba644c731ab54abd6f7c3568"
    repository = "https://gitlab.com/Dimbreath/turnbasedgamedata"
    wiki = "https://wiki.biligame.com/sr/index.php?title=%E4%BA%8B%E4%BB%B6%E4%B8%80%E8%A7%88&oldid=92645"
    data = {
        "Sources": [
            [1, "du.source.color.handbook", repository, revision, "4.4", "2026-09-08",
             "ExcelOutput/RogueTournHandBookEvent.json; EventHandbookID=108",
             "15e0e448186b799687a1ef4582e96b40da43888fee5a0e31278eb6cbc29106b9",
             "ExactStructured", "Git blob digest; current handbook identity, not option mechanics."],
            [2, "du.source.color.variant", repository, revision, "4.4", "2026-09-08",
             "ExcelOutput/RogueTournNPC.json; RogueNPCID=722601",
             "e3ef2da06848f45704310e1a35ae06390febf7d4e5a6027591912929b4e4c478",
             "ExactStructured", "Git blob digest; published graph absent at pinned revision."],
            [3, "du.source.color.public-options", wiki, "92645", "4.4", "2026-09-08",
             "A Dash of Color / 一抹色彩; three alternative outcomes",
             "c0a1e384cd77623e8f95ebcc0ac5491e73bf946890e309fde6dc40bee630adc1",
             "ObservedCommunity", "Canonical event-list HTML digest. Values reported by community; application to 4.4 is policy, not observed parity."],
        ],
        "Policies": [[
            1, "du.policy.color.current-options", "4.4", "VersionedProjectPolicy",
            "UniformUnownedCurrentCatalog", "DisableChoice",
            "Apply public 200/2/2 outcomes to the current variant; zero fragment costs. For random rewards, filter current Common/Rare identities, exclude owned, choose one canonical current state per Curio and sample identities uniformly without replacement. Insufficient pool disables that choice; fragment choice remains available.",
            "Replace individual fields when released current-variant graph or reproducible observations establish binding, costs, rarity mapping, pool/state eligibility, weights or exhaustion behavior.",
            "1|2|3",
        ]],
        "Occurrences": [[
            1, "du.occurrence.a-dash-of-color", "divergent-universe.occurrence.108",
            "divergent-universe.occurrence-variant.722601", 1, "1|2",
        ]],
        "Choices": [
            [1, "du.choice.color.pale-gold", 1, 1, "Pale gold", "淡金色", 0, 3],
            [2, "du.choice.color.blood-red", 1, 2, "Blood red", "血红色", 0, 3],
            [3, "du.choice.color.dark-gray", 1, 3, "Dark gray", "深灰色", 0, 3],
        ],
        "Outcomes": [
            [1, "du.outcome.color.fragments", 1, 1, "Fragments", 200, None, None, 3],
            [2, "du.outcome.color.curios", 2, 1, "Curios", 2, 1, 2, 3],
            [3, "du.outcome.color.blessings", 3, 1, "Blessings", 2, 1, 2, 3],
        ],
    }
    for row_id, key, locator, digest, note in [
        (4, "curio-mode-copy", "ExcelOutput/RogueTournMiracle.json; MiracleID=9028,9053,9167,9195,9196,9197; TournMode=Tourn3",
         "176b17030bdf212920d8a59142cf982916349ee986b574b3e2d3bbdeff81c438", "Exact current mode-copy to effect bindings; no shared-name inference."),
        (5, "curio-acquisition-effects", "ExcelOutput/RogueMiracleEffect.json; MiracleEffectID=2028,2053,2167,2195,2196,2197; ParamList[0]",
         "fd40b0b35b2735875356681597a532537fb4a98fbcddded501e8d4b4e3eede35", "Exact effect parameters and description text hashes; full lifecycle is not implemented by an acquisition row."),
        (6, "curio-acquisition-en", "TextMap/TextMapEN.json; hashes=2404701295559821740,16761194725868865488,259359941151777850,1741960291225087902,7179292461667666903,2000972934457142612",
         "afc6d0ff9eeb20d09042e61c311e60d4ed56e69757f1ac42466ca4840e6ed789", "Released EN text establishes immediate fragment gain; scheduling and rounding remain policy."),
    ]:
        data["Sources"].append([row_id, f"du.source.{key}", repository, revision, "4.4", "2026-09-08", locator, digest, "ExactStructured", note])
    data["CurioAcquisitions"] = [
        [ordinal, f"du.curio-acquisition.{state}", f"divergent-universe.curio-state.{state}", str(effect), kind, 1, amount,
         "StableStateOrderFloorBeforeEachGrant", summary_en, summary_zh,
         "Execute the reviewed immediate grant once after inventory insertion, in stable state-key order within one accepted batch. Fractional grants read the balance immediately before that grant and floor nonnegative fractional fragments. Other acquisition and gain-modifier effects remain implementation gaps, not no-effect policies.",
         "Replace batch ordering, balance snapshot or rounding independently when a released execution graph or reproducible observation establishes them; preserve exact amounts and atomic rollback.", "4|5|6", None]
        for ordinal, (state, effect, kind, amount, summary_en, summary_zh) in enumerate([
            (9028, 2028, "FixedFragments", "300", "Gain 300 fragments on acquisition; later domain debt remains separate.", "获取时获得 300 碎片；后续区域债务另行处理。"),
            (9053, 2053, "BalanceFraction", "0.4", "Gain fragments equal to 40% of the current balance on acquisition.", "获取时获得当前持有碎片数量 40% 的碎片。"),
            (9167, 2167, "FixedFragments", "500", "Gain 500 fragments on acquisition; domain expiry and penalties remain separate.", "获取时获得 500 碎片；区域失效与惩罚另行处理。"),
            (9195, 2195, "FixedFragments", "150", "Green Miracle (Dormant) grants 150 fragments on acquisition.", "绿奇迹（休眠）在获取时获得 150 碎片。"),
            (9196, 2196, "FixedFragments", "300", "Green Miracle (Awakened) grants 300 fragments on acquisition.", "绿奇迹（觉醒）在获取时获得 300 碎片。"),
            (9197, 2197, "FixedFragments", "600", "Green Miracle (Exalted) grants 600 fragments on acquisition.", "绿奇迹（升华）在获取时获得 600 碎片。"),
        ], 1)
    ]
    for row_id, source_key, locator, digest in [
        (7, "wax-mode-copies", "ExcelOutput/RogueTournMiracle.json; MiracleID=9043..9049,9147,9187; TournMode=Tourn3", "176b17030bdf212920d8a59142cf982916349ee986b574b3e2d3bbdeff81c438"),
        (8, "wax-effects", "ExcelOutput/RogueMiracleEffect.json; MiracleEffectID=2043..2049,2147,2187; ParamList[0]", "fd40b0b35b2735875356681597a532537fb4a98fbcddded501e8d4b4e3eede35"),
        (9, "wax-text", "TextMap/TextMapEN.json; hashes=12222930147806732735,16753531798159570738,12337916719786405821,15397308621934325114,3468534171698544820,2964235700804576905,12725242702040094842,15613011178341524346,12763520219546358679", "afc6d0ff9eeb20d09042e61c311e60d4ed56e69757f1ac42466ca4840e6ed789"),
        (10, "wax-path-types", "ExcelOutput/RogueTournBuffType.json; RogueBuffType=121,122,124,125,126,127,128,129", "f9163c3414bf357c2a66d3a8c26b94e0ffda0a520a233af74d764f09a406c7f5"),
        (11, "wax-active-paths", "ExcelOutput/RogueTournUseBuffType.json; TournMode=Tourn3; UseBuffTypeList", "506229616f84b385e61608d39541c1da0876e6fbe1b7775bc87f84f0849784f7"),
    ]:
        data["Sources"].append([row_id, f"du.source.{source_key}", repository, revision, "4.4", "2026-09-08", locator, digest, "ExactStructured", "Released state/effect/count/path evidence; hidden pool weights, ordering and exhaustion remain explicit policy."])
    for ordinal, (state, effect, paths, name, name_zh) in enumerate([
        (9043, 2043, "126", "Elation", "欢愉"),
        (9044, 2044, "124", "The Hunt", "巡猎"),
        (9045, 2045, "125", "Destruction", "毁灭"),
        (9046, 2046, "121", "Remembrance", "记忆"),
        (9047, 2047, "122", "Nihility", "虚无"),
        (9048, 2048, "127", "Propagation", "繁育"),
        (9049, 2049, "128", "Erudition", "智识"),
        (9147, 2147, "129", "Harmony", "同谐"),
        (9187, 2187, "121|122|124|125|126|127|128|129", "each current active Path", "每个当前活动命途"),
    ], 7):
        data["CurioAcquisitions"].append([
            ordinal, f"du.curio-acquisition.{state}", f"divergent-universe.curio-state.{state}", str(effect), "PathBlessings", 1, "1",
            "StableStateOrderUniformUnownedPathRejectExhaustion",
            f"On acquisition, grant one random Blessing of {name}; later offer weights or Equation triggers remain separate.",
            f"获取时获得{ name_zh }的随机祝福各一个；后续候选权重或方程触发另行处理。",
            "In stable state-key and path-key order, sample current path-matching unowned Blessing identities uniformly without replacement across the batch. All rarities are eligible. Before drawing, prove that all mandatory Path and rarity grants have a complete distinct assignment; each draw excludes candidates that would prevent the remaining grants. This conditions mixed-pool draws on feasibility, not on uniform complete assignments. Reject acquisition without RNG/state change if mandatory rewards cannot be completed or a Blessing offer/progress update is pending. Reward pools exclude such unavailable Curios. Coalesce all mandatory Blessing grants after the batch's stable-order fragment effects, then refresh Equation progress once before authored expansion rewards and room completion. Expansion rewards use their separate policy and merge allowance changes into final Curio holdings. Other acquisition effects remain implementation gaps, not no-effect policies.",
            "Replace hidden membership, rarity weights, batch/path ordering, exhaustion or pending-offer scheduling independently when released execution data or reproducible observations establish them; retain exact path/count and atomic rollback.",
            "7|8|9|10|11", paths,
        ])
    for row_id, source_key, locator, digest in [
        (12, "fragment-gain-states", "ExcelOutput/RogueTournMiracle.json; MiracleID=9055,9070,9079,9159; TournMode=Tourn3", "176b17030bdf212920d8a59142cf982916349ee986b574b3e2d3bbdeff81c438"),
        (13, "fragment-gain-effects", "ExcelOutput/RogueMiracleEffect.json; MiracleEffectID=2055,2070,2079,2159; ParamList[0]", "fd40b0b35b2735875356681597a532537fb4a98fbcddded501e8d4b4e3eede35"),
        (14, "fragment-gain-text", "TextMap/TextMapEN.json; hashes=9113138745615695677,7205940712072507551,3429135723207096768,4340701535028103385", "afc6d0ff9eeb20d09042e61c311e60d4ed56e69757f1ac42466ca4840e6ed789"),
    ]:
        data["Sources"].append([row_id, f"du.source.{source_key}", repository, revision, "4.4", "2026-09-08", locator, digest, "ExactStructured", "Exact state/effect/text facts. Expiry, battle-Blessing suppression and price penalties have separate runtime owners; the gain component alone does not execute them."])
    data["CurioFragmentGains"] = [
        [ordinal, f"du.curio-fragment-gain.{state}", f"divergent-universe.curio-state.{state}", str(effect), 1, fraction,
         "ActiveStateAdditiveOriginalBaseFloorEachBonus",
         f"While active, increase Cosmic Fragment gains by {percent}%; other effects remain separate.",
         f"生效期间使宇宙碎片获取量提高 {percent}%；其他效果另行处理。",
         "Snapshot the nonnegative original gain once. Credit it, then add each active state bonus independently in stable state-key order, flooring each bonus from that original base. Bonuses never recursively trigger bonuses. Acquisition batches insert all accepted holdings before gains. Destroyed/replaced states stop contributing immediately; repairs restore contribution. All operations and overflow rejection share the owning transaction. Expiry and price penalties remain implementation gaps, not no-effect policies. Normal-battle Blessing suppression is authored separately.",
         "Replace additive versus multiplicative stacking, per-bonus rounding, acquisition activation timing and refund applicability independently when released execution data or reproducible observations establish them; exact rates remain source-backed. Current policy includes every positive fragment credit and excludes spends/resets.", "12|13|14"]
        for ordinal, (state, effect, fraction, percent) in enumerate([
            (9055, 2055, "0.5", 50), (9070, 2070, "0.5", 50),
            (9079, 2079, "0.3", 30), (9159, 2159, "0.3", 30),
        ], 1)
    ]
    data["Sources"].append([
        15, "du.source.battle-blessing-catalog", repository, revision, "4.4", "2026-09-08",
        "ExcelOutput/RogueTournBuff.json; current reference catalog base-level identities and RogueBuffCategory",
        "4bcf0b062941400a7aa4a530b57a5e9e42e4a0a1d8b4a67e90c7886214c27cb7",
        "ExactStructured", "Exact identity/category records, not evidence for normal-battle pool membership, offer width or weights.",
    ])
    data["BattleBlessings"] = [[
        1, "du.battle-blessings.normal", "UniformUnownedSingleSelectionAvailableSubset", 3, 1, 2,
        "divergent-universe.curio-state.9055", "2055",
        "After a normal battle victory, choose one offered unowned base Blessing. Active reviewed suppression prevents the offer.",
        "普通战斗胜利后，从候选中选择一个未拥有的基础祝福；已核实的生效奇物限制会阻止此次奖励。",
        "VersionedProjectPolicy: apply to the explicit baseline normal-battle route in both run families. Use uniform base weights for distinct unowned current Common/Rare identities in stable identity order, apply separately authored active CurioBattleWeights bonuses, and sample up to three without replacement. Present candidates in stable numeric state-key order and accept exactly one. An exhausted or suppressed pool advances without a Blessing. Suppression and weights snapshot active state at verified settlement. Fragment drops, extra selections, base rarity/path weighting, rerolls, elite/boss rewards and other acquisition triggers remain pending, not zero-effect policies.",
        "Replace individual width, membership, rarity/Path weights, exhausted-pool behavior, suppression snapshot timing and encounter binding when released current execution data or reproducible observations establish them. Do not promote this chosen current-catalog pool to exact membership.",
        "12|13|14|15",
    ]]
    data["CurioBattleWeights"] = [
        [ordinal, f"du.curio-battle-weight.{state}", f"divergent-universe.curio-state.{state}", str(effect), path, 3,
         "ActiveStateAdditiveCandidateWeight",
         f"While active, favor {name} Blessings in normal-battle victory offers; numeric weight is project policy.",
         f"生效期间提高普通战斗胜利候选中{zh}祝福的抽取权重；具体数值为项目策略。",
         "VersionedProjectPolicy: at verified settlement, add three to each matching unowned candidate's base weight of one for every active bound state. Use checked integer addition and shared weighted sampling without replacement. Different paths retain independent bonuses; no guaranteed slot or extra RNG stream. Destroyed/replaced states contribute nothing and repair restores the bonus. Existing offers retain their snapshot. This four-to-one candidate-weight ratio is a low-confidence approximation of increased chance, not a measured fourfold Path probability. Other reward sources and Trailblaze Wax Equation triggers are separate.",
         "Replace magnitude, additive stacking, base rarity weights, eligible reward boundaries and snapshot timing independently when released execution data or reproducible observations establish them. Alternatives include multiplicative weights and guaranteed Path slots; positive additive weights preserve increased chance without inventing a guarantee.",
         "7|8|9|10|11|15"]
        for ordinal, (state, effect, path, name, zh) in enumerate([
            (9043, 2043, "126", "Elation", "欢愉"),
            (9044, 2044, "124", "The Hunt", "巡猎"),
            (9045, 2045, "125", "Destruction", "毁灭"),
            (9046, 2046, "121", "Remembrance", "记忆"),
            (9047, 2047, "122", "Nihility", "虚无"),
            (9048, 2048, "127", "Propagation", "繁育"),
            (9049, 2049, "128", "Erudition", "智识"),
            (9147, 2147, "129", "Harmony", "同谐"),
        ], 1)
    ]
    data["EquationGrants"] = [[
        1, "du.equation-grant.trailblaze", "divergent-universe.curio-state.9187", "2187", 2, 3, 3, 1,
        "UniformMissingRecipeAvailableSubsetOnceLogicalDomain",
        "While active, obtaining an Equation grants up to three unowned Blessings still needed by its recipe, once per logical Domain visit.",
        "生效期间获得方程时，补充最多三个该方程仍缺少的未拥有祝福；每个逻辑区域访问实例限一次。",
        "VersionedProjectPolicy: normal acquisition and replacement both count as obtaining the new Equation. Sequentially sample uniformly from current unowned identities that contribute to a still-deficient main or sub Path, recomputing deficits after each grant, across all rarities. Stop at the authored count or when none remain. Consume the once budget only when at least one Blessing is granted; already satisfied/exhausted recipes consume neither budget nor RNG. Snapshot active Curio status before acquisition. Track the existing logical Node address and visit, not physical battle/reward nodes. Destroy, repair or reacquisition cannot reset this receipt. All holdings, progress, receipt, offer consumption and RNG commit atomically. Pending Blessing offers reject when a grant would execute; no recursive acquisition triggers are implied.",
        "Replace replacement-trigger inclusion, candidate membership, rarity/Path weights, deficit scheduling, subset/empty-budget behavior, pending-offer ordering or logical-to-physical Domain binding independently when released execution topology or reproducible observations establish them. Alternatives include all matching Paths even if satisfied, main-first selection, and consuming an empty trigger. Current policy preserves useful missing-recipe grants without inventing duplicates or extra counts; parity confidence is low for these fields.",
        "7|9|10|11|15|16",
    ]]
    data["Sources"].append([
        16, "du.source.trailblaze-equation-counts", repository, revision, "4.4", "2026-09-09",
        "ExcelOutput/RogueMiracleEffect.json; MiracleEffectID=2187; ParamList[1]=3, ParamList[2]=1; MiracleDesc.Hash=12763520219546358679",
        "fd40b0b35b2735875356681597a532537fb4a98fbcddded501e8d4b4e3eede35", "ExactStructured",
        "Exact trigger count and per-Domain limit; missing-recipe sampling, empty-trigger handling, replacement inclusion and logical Domain mapping remain explicit policy.",
    ])
    data["Sources"].append([
        17, "du.source.initial-equation-catalog", repository, revision, "4.4", "2026-09-09",
        "ExcelOutput/RogueTournFormula.json; current reference Equation identities and categories",
        "87778b3d411554657e9609d2e364395bcbc55ba38f8b15e1d250e475e81f2598", "ExactStructured",
        "Exact identity/category records only. Initial placement, pool membership, offer width, weights and application to both run families are project policy, not established by RandomID or by a different iteration's public guide.",
    ])
    data["InitialEquations"] = [[
        1, "du.initial-equations.current", "UniformCurrentCategorySingleSelection", 3, "Epic",
        "Choose one initial Equation from a uniformly sampled current Epic-category pool before the baseline event.",
        "基线事件前，从当前 Epic 类别均匀抽取的候选中选择一个初始方程。",
        "VersionedProjectPolicy: both normal baseline run families start with a required three-option Equation selection. Use distinct current Epic-category identities, stable identity ordering, unit weights and shared Reward RNG without replacement. Acquire exactly one through the ordinary ownership/progress/acquisition-grant transaction, then enter the existing baseline event. No reroll, hidden candidate, internal offer overwrite or raw graph-choice bypass is permitted. This does not implement missing mask/Titan selection, Great Blessings, initial Blessing batches or complete room placement; they remain gaps, not zero-effect policies. Confidence in placement, category mapping, width, pool membership and weights is low.",
        "Replace each field independently when pinned released execution data or reproducible observations bind initial selection to this exact current profile. Alternatives include fixed entry-specific pools, other rarity categories, weighted paths and scheduling after entry-specific boons. Current policy provides an explicit playable choice without claiming candidate membership or full starting-loadout parity.",
        "17",
    ]]
    data["Sources"].append([
        18, "du.source.layer-battle-route", repository, revision, "4.4", "2026-09-09",
        "ExcelOutput/RogueTournArea.json; current reference area-to-ordered-layer joins",
        "756510177a468130464c27ef382d9ab33a7098dc993022cd8647366cc7c46db8", "ExactStructured",
        "Exact area/layer identities and order only; per-layer battle count, room membership, encounter placement and boss boundaries remain project policy.",
    ])
    data["BattleRoutes"] = [[
        1, "du.battle-route.logical-layers", "OneBattlePerLayerWithDomainChoices", 1,
        "Each logical layer executes one proxy battle; the first domain is Combat and later domains are selected through the authored public choices.",
        "每个逻辑层执行一场代理战斗；首个区域为普通战斗，后续区域通过配置的公开选项选择。",
        "VersionedProjectPolicy: both baseline families follow the exact ordered area/layer joins, place one normal battle at each logical checkpoint, and use the separately authored EncounterPools candidate/proxy policy. Battle results carry HP, energy, life and presence through the shared ledger; every victory generates the authored normal-battle Blessing offer and every loss/fault terminates without victory rewards. Separate physical battle/reward nodes retain their owning logical layer scopes. No automatic healing, retry, extra room, elite reward or boss program is implied. Room positions, branch counts, room-specific encounter pools, elite/boss placement and full content reachability remain implementation gaps; this is not complete gameplay parity.",
        "Replace the low-confidence count, normal-battle classification, checkpoint placement, encounter assignment and room/plane scope mapping independently when released current-profile room graphs or reproducible observations establish them. Alternatives include several rooms per layer and mixed combat/event/service branches. Current policy exercises real multi-battle carry and settlement without claiming an unpublished room pool or promoting candidate enemies to exact bindings.",
        "18",
    ]]
    for source_id, name, locator, digest in [
        (19, "baseline-encounter-group", "ExcelOutput/RogueMonsterGroup.json; RogueMonsterGroupID=300202; RogueMonsterListAndWeight", "8abe52c1f9b5fd42eeefaee63ff8bc59c8dd8acc0f18502b5548b353416f2171"),
        (20, "baseline-encounter-stages", "ExcelOutput/RogueMonster.json; RogueMonsterID=3002081,3002111,3002161,3002181,3002191; EventID", "52f235fecb6aa99d3fe01d4f1ff8c42ac3b7c8fbd53fb77a7bbdae140e056f32"),
    ]:
        data["Sources"].append([source_id, f"du.source.{name}", repository, revision, "4.4", "2026-09-09", locator, digest, "ExactStructured", "Exact group-to-monster-to-stage joins only. Weekly-display candidates do not prove normal-room membership or enabled current-profile selection."])
    data["EncounterPools"] = [[
        1, "du.encounter-pool.baseline", "FixedFirstUniformLaterLayersWithReplacement",
        "divergent-universe.encounter-group.300202", "83002081",
        "83002081|83002111|83002161|83002181|83002191",
        "Use the authored first encounter, then independently sample one candidate for each later logical layer.",
        "首层使用配置的固定遭遇，后续每个逻辑层独立抽取一个候选遭遇。",
        "VersionedProjectPolicy: both baseline families use the explicitly listed weekly-display candidate stages as normal proxy encounters. Fix the first layer to the authored first stage; later layers independently select one stage with unit integer weights in stable stage order, with replacement across layers, using shared Encounter RNG when the layer offer is created. Present only the selected encounter and bind assembly to that offer. Group/stage joins are source-backed; normal-room placement, first-stage choice, cross-layer reuse and eligibility are low-confidence policy. This does not enable the weekly boss selector or implement original elite/boss programs.",
        "Replace group membership, room/layer binding, first-stage choice, weights and replacement independently when released current-profile selectors or reproducible observations establish them. Alternatives include room-specific pools, weighted sampling, no-repeat rules and first-layer randomness. Preserve atomic offer generation and assembly authentication; do not promote display candidates or proxy mechanics to exact parity.",
        "19|20",
    ]]
    for source_id, name, locator, digest, note in [
        (21, "green-battle-state", "ExcelOutput/RogueTournMiracle.json; MiracleID=9195; TournMode=Tourn3; MiracleEffectID=2195; HandbookMiracleID=9155", "176b17030bdf212920d8a59142cf982916349ee986b574b3e2d3bbdeff81c438", "Exact current owned mode-copy binding; unbound upgrade states 9196/9197 remain outside this executable reward row."),
        (22, "green-battle-effect", "ExcelOutput/RogueMiracleEffect.json; MiracleEffectID=2195; ParamList[1]=8", "fd40b0b35b2735875356681597a532537fb4a98fbcddded501e8d4b4e3eede35", "Exact per-full-HP-character amount. Does not establish ordering against healing, global gain modifiers or other battle reward programs."),
        (23, "green-battle-text", "TextMap/TextMapEN.json; Hash=1741960291225087902", "afc6d0ff9eeb20d09042e61c311e60d4ed56e69757f1ac42466ca4840e6ed789", "Released description establishes the victory/full-HP trigger. Occurrence upgrades have separate transition and public-event policy bindings."),
        (24, "baseline-stage-classification", "ExcelOutput/StageConfig.json; StageID=83002081,83002111,83002161,83002181,83002191; StageConfigData._IsEliteBattle=1; StageAbility_RogueEliteLevel_Enhance", "89e7dcd4824b81b080f985930142ddd903db4f31fc41667e898c0022a85ced7f", "Exact elite-stage markers. Baseline normal proxy classification is a deliberate non-parity policy, not absent or uncertain source classification; original elite/boss execution remains pending."),
    ]:
        data["Sources"].append([source_id, f"du.source.{name}", repository, revision, "4.4", "2026-09-09", locator, digest, "ExactStructured", note])
    data["EncounterPools"][0][-1] = "19|20|24"
    data["CurioBattleGrants"] = [[
        1, "du.curio-battle-grant.dormant-green", "divergent-universe.curio-state.9195", "2195", 2, 8,
        "FullHpPresentRosterAfterCarry",
        "After victory, active Green Miracle (Dormant) grants eight fragments for each present, living character at full HP.",
        "战斗胜利后，生效中的绿奇迹（休眠）为每名在场、存活且满生命值的角色提供八枚碎片。",
        "VersionedProjectPolicy: snapshot the verified post-battle participant carry ledger before any additional mode healing. Count current locked participants only, with Alive/Present status and current HP equal to positive maximum HP. Active Curio status is checked at settlement; absent, destroyed or replaced states grant nothing. Credit one aggregate amount through the ordinary fragment-gain pipeline before the Blessing offer is published. Global active gain bonuses apply once to this aggregate, with their existing independent floor rules. Zero eligible characters produces no credit or RNG. Result validation, carry, credit, offer generation and graph advance are one transaction. Base battle drops, battle-only gain multipliers/suppression, original Occurrence placement and other Divine Treasure producers remain gaps, not zero-effect policies.",
        "Replace snapshot timing, off-field/defeated participant eligibility, aggregation versus per-character credit, and interaction ordering independently when released execution data or reproducible observations establish them. Exact amount and victory/full-HP trigger remain source-backed. Alternatives include pre-carry HP, post-healing HP and separate per-character rounding; current ordering uses the verified shared settlement boundary and never grants on loss/fault or duplicate results. Confidence in hidden scheduling/stacking is low.",
        "21|22|23",
    ]]
    for source_id, name, locator, digest, note in [
        (25, "green-display-names", "ExcelOutput/RogueMiracleDisplay.json; MiracleDisplayID=256,257,258; name hashes=4410465334307410196,16129282454969858300,10527537750345113790", "5937214885a6f3c2f6fcc65870ab7bccd4e4493cffd682dbda75e63ac50ce739", "Exact display/name joins distinguish Dormant, Awakened and Exalted. Name similarity alone does not prove an upgrade edge or handbook membership."),
        (26, "divine-treasures-handbook", "ExcelOutput/RogueTournHandBookEvent.json; EventHandbookID=165,166; NPC=627401,727401; progress=1,2", "15e0e448186b799687a1ef4582e96b40da43888fee5a0e31278eb6cbc29106b9", "Exact event/variant/progress identities. Referenced NPC graphs are absent at the pinned revision; costs, probability and execution scheduling are not established."),
    ]:
        data["Sources"].append([source_id, f"du.source.{name}", repository, revision, "4.4", "2026-09-10", locator, digest, "ExactStructured", note])
    for source_id, name, title, page_revision, digest, note in [
        (27, "green-awakened-public", "%E7%BB%BF%E5%A5%87%E8%BF%B9%EF%BC%88%E8%A7%89%E9%86%92%EF%BC%89", "77066", "d8b52c199ccb711b3b7b286fedec8ab5a3b94c533f8aa4bfc2205bb8b4990184", "Awakened variant is related to Divine Treasures II. Community iteration binding is not exact current-profile reachability."),
        (28, "green-exalted-public", "%E7%BB%BF%E5%A5%87%E8%BF%B9%EF%BC%88%E5%8D%87%E5%8D%8E%EF%BC%89", "77072", "422043bfd73b1900ea25561f22e04340161eabff5fc9113276dfabf7e01e4a9a", "Exalted variant is related to Divine Treasures III. Page reports 64 per full-HP character, conflicting with pinned 4.4 effect 2197 value 32; only event association is used, never the conflicting amount."),
    ]:
        data["Sources"].append([source_id, f"du.source.{name}", f"https://wiki.biligame.com/sr/{title}", page_revision, "4.4", "2026-09-10", "Curio detail table; related Occurrence; fetched HTML bytes", digest, "ObservedCommunity", note])
    data["Sources"].append([
        29, "du.source.green-upgrade-victory-amounts", repository, revision, "4.4", "2026-09-10",
        "ExcelOutput/RogueMiracleEffect.json; MiracleEffectID=2196,2197; ParamList[1]=16,32; description hashes=7179292461667666903,2000972934457142612",
        "fd40b0b35b2735875356681597a532537fb4a98fbcddded501e8d4b4e3eede35", "ExactStructured",
        "Exact upgraded full-HP victory amounts. Pinned 4.4 value 32 takes precedence over the conflicting community Exalted value 64.",
    ])
    data["CurioEvolutions"] = [
        [ordinal, f"du.curio-evolution.{before}.{after}", f"divergent-universe.curio-state.{before}", f"divergent-universe.curio-state.{after}", "divergent-universe.curio.9155", f"divergent-universe.occurrence.{event}", "ActiveSameOwnerResetAndAcquire",
         f"Evolve Green Miracle from {before} to {after}, preserving one ownership identity and executing the new state's acquisition grant.",
         f"将绿奇迹从形态 {before} 升级为 {after}，保留单一持有身份并执行新形态的获取奖励。",
         "VersionedProjectPolicy: bind the independently reported Dormant/Awakened/Exalted sequence to the current named mode-copy states. Missing handbook IDs on targets remain absent in the reference catalog; add a separate runtime ownership alias only for this accepted evolution. Require the active predecessor, remove its lifecycle/counters, insert the new active state with declared initial charges and zero activations, then execute the full new acquisition grant atomically. No delta subtraction, inherited counters, ordinary-pool insertion, direct acquisition or generic replacement into evolution-only targets. This trusted boundary itself does not authorize a player choice, place an event, charge its price, roll its probability or grant its other rewards; those producer rules are authored separately in EvolutionEvents and EvolutionOptions.",
         "Replace current-profile edge/owner binding, active-only eligibility, counter reset and full-grant-on-evolution independently when released execution graphs or reproducible observations establish them. Alternatives include standalone ownership identities, retained counters, destroyed-state evolution and difference-only credits. The selected policy preserves one Curio and source-backed new-state effects without inventing extra ordinary rewards; parity confidence is low for hidden transition details.",
         f"4|5|6|25|26|{source}"]
        for ordinal, (before, after, event, source) in enumerate([(9195, 9196, 165, 27), (9196, 9197, 166, 28)], 1)
    ]
    for ordinal, state, effect, amount in [(2, 9196, 2196, 16), (3, 9197, 2197, 32)]:
        row = list(data["CurioBattleGrants"][0])
        row[:6] = [ordinal, f"du.curio-battle-grant.{state}", f"divergent-universe.curio-state.{state}", str(effect), 2, amount]
        row[7:9] = [f"After victory, active evolved Green Miracle grants {amount} fragments per qualifying full-HP character.", f"战斗胜利后，生效中的升级绿奇迹为每名符合条件的满血角色提供 {amount} 碎片。"]
        row[-1] = "4|6|29"
        data["CurioBattleGrants"].append(row)
    append_evolution_events(data)
    for row in data["CurioAcquisitions"]:
        row.extend([None, None])
    append_sage_acquisitions(data)
    append_sage_victory_blessings(data)
    append_domain_choices(data)
    append_battle_fragments(data)
    append_curio_expiries(data)
    append_curio_battle_stats(data)
    append_curio_battle_reactions(data)
    append_tawot_services(data)
    append_curio_domain_grants(data)
    append_tawot_victory(data)
    append_equation_expansion_rewards(data)
    append_domain_layout(data)
    append_domain_decks(data)
    return data


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError(f"refusing to overwrite {args.output}")
    template = args.root / "config/divergent-universe-decisions-generated/templates/DivergentUniverseDecisions.xlsx"
    workbook = load_workbook(template)
    data = rows()
    if workbook.sheetnames != list(data):
        raise ValueError("template sheet order differs from authoring contract")
    for name, records in data.items():
        sheet = workbook[name]
        if sheet.max_row != 7:
            raise ValueError(f"{name}: template must contain exactly seven metadata rows")
        columns = [cell.column for cell in sheet[3] if cell.value not in (None, "#field")]
        for row_index, record in enumerate(records, 8):
            if len(record) != len(columns):
                raise ValueError(f"{name}: typed column count differs")
            for column, value in zip(columns, record, strict=True):
                sheet.cell(row_index, column, value)
        sheet.freeze_panes = "A8"
        sheet.auto_filter.ref = f"A3:{sheet.cell(3, sheet.max_column).column_letter}{sheet.max_row}"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    workbook.close()
    print(f"Authored {sum(map(len, data.values()))} decision rows using openpyxl.")


if __name__ == "__main__":
    main()
