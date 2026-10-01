"""Clean-target openpyxl authoring for the dedicated damage input capability.

These unbound ProjectPolicy probes do not represent released character kits.
Sora templates own headers; all existing live cells must remain unchanged.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "config" / "data"
STAT_TABLES = ("ModifierDefinition", "SnapshotCapture")


def encoded(value: dict[str, object]) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def records() -> dict[str, list[dict[str, object]]]:
    literals = {
        970101: ("resolved-base", "100.9"),
        970102: ("original", "0.8"),
        970103: ("merrymaking", "-0.5"),
        970104: ("target-resistance", "-0.2"),
        970105: ("penetration", "0.1"),
        970106: ("resistance-minimum", "-1"),
        970107: ("resistance-maximum", "0.9"),
        970108: ("unbroken", "0.9"),
        970109: ("one", "1"),
        970112: ("property-addition", "0.5"),
    }
    expressions = [{
        "id": identifier,
        "stable_key": f"policy.probe.elation.expr.{name}",
        "result_kind": "Scalar",
        "node": encoded({"type": "ScalarLiteral", "value_decimal": value}),
    } for identifier, (name, value) in literals.items()]
    for identifier, name, kind, node in [
        (970110, "owner-property", "Scalar", {
            "type": "QueryStat", "subject_selector_id": 24051,
            "stat": "Elation", "formula_purpose": "Stat",
        }),
        (970111, "resolved-meter", "Scalar", {
            "type": "CheckedBinary", "operator": "CheckedAdd",
            "left_expression_id": 970109, "right_expression_id": 970110,
            "rounding": "NearestTiesEven",
        }),
        (970113, "non-scalar-negative", "Integer", {
            "type": "IntegerLiteral", "value": 1,
        }),
    ]:
        expressions.append({
            "id": identifier,
            "stable_key": f"policy.probe.elation.expr.{name}",
            "result_kind": kind, "node": encoded(node),
        })
    return {
        "ValueExpression": sorted(expressions, key=lambda record: record["id"]),
        "Operation": [{
            "id": 970201,
            "stable_key": "policy.probe.elation.operation.explicit-self-damage",
            "domain": "Battle", "target_selector_id": 24051,
            "condition_id": None, "empty_target_policy": "Fault",
            "snapshot_boundary": "Dynamic", "fault_policy": "Rollback",
            "payload": encoded({
                "type": "ElationDamage",
                "base_damage_expression_id": 970101,
                "original_multiplier_expression_id": 970102,
                "meter_multiplier_expression_id": 970111,
                "merrymaking_expression_id": 970103,
                "target_resistance_expression_id": 970104,
                "penetration_expression_id": 970105,
                "resistance_minimum_expression_id": 970106,
                "resistance_maximum_expression_id": 970107,
                "unbroken_multiplier_expression_id": 970108,
                "element": "Fire", "can_crit": False,
            }),
        }],
    }


def cell_projection(sheet) -> dict[tuple[int, int], tuple[object, ...]]:
    return {(cell.row, cell.column): (
        cell.value, cell.data_type,
        tuple(cell._style) if cell.has_style else None,
        cell.comment.text if cell.comment else None,
        cell.hyperlink.target if cell.hyperlink else None,
    ) for row in sheet for cell in row}


def stat_dropdown(sheet, validations=None):
    columns = [cell.column for cell in sheet[3] if cell.value == "stat"]
    if len(columns) != 1:
        raise ValueError("Sora stat header is missing or ambiguous")
    column = columns[0]
    coordinate = f"{get_column_letter(column)}8"
    collection = sheet.data_validations if validations is None else validations
    matches = [rule for rule in collection.dataValidation
               if rule.type == "list" and coordinate in rule.sqref]
    if len(matches) != 1 or any(
        area.min_col != column or area.max_col != column
        for rule in matches for area in rule.sqref.ranges
    ):
        raise ValueError("stat dropdown ownership is missing or ambiguous")
    return matches[0]


def stat_formula(table: str) -> str:
    template = load_workbook(ROOT / "config" / "generated" / "templates" / f"{table}.xlsx")
    try:
        return stat_dropdown(template.active).formula1
    finally:
        template.close()


def check_rows(sheet, expected: list[dict[str, object]], append: bool) -> None:
    fields = [cell.value for cell in sheet[3][1:] if cell.value]
    id_column = fields.index("id") + 2
    for record in expected:
        if set(record) != set(fields):
            raise ValueError("authoring record differs from Sora-owned headers")
        matches = [row for row in range(8, sheet.max_row + 1)
                   if str(sheet.cell(row, id_column).value) == str(record["id"])]
        if len(matches) > 1:
            raise ValueError(f"duplicate owned ID {record['id']}")
        if matches:
            current = {field: sheet.cell(matches[0], column).value
                       for column, field in enumerate(fields, start=2)}
            if current != record:
                raise ValueError(f"owned ID {record['id']} conflicts with designer data")
        elif append:
            row = max(8, sheet.max_row + 1)
            for column, field in enumerate(fields, start=2):
                sheet.cell(row, column).value = record[field]
        else:
            raise ValueError(f"missing owned ID {record['id']}")


def write_clean(output: Path) -> None:
    output = output.resolve()
    if output.exists() or output == DATA or DATA in output.parents:
        raise ValueError("authoring requires a new non-production output directory")
    output.mkdir(parents=True)
    owned_records = records()
    for table in [*owned_records, *STAT_TABLES]:
        expected = owned_records.get(table, [])
        source = DATA / f"{table}.xlsx"
        source_digest = hashlib.sha256(source.read_bytes()).hexdigest()
        workbook = load_workbook(source)
        try:
            sheet = workbook.active
            before = cell_projection(sheet)
            validations = deepcopy(sheet.data_validations)
            if expected:
                check_rows(sheet, expected, append=True)
            if table in STAT_TABLES:
                # Only this schema-owned enum choice formula changes. Preserve
                # every source range, validation setting and other dropdown.
                stat_dropdown(sheet).formula1 = stat_formula(table)
                stat_dropdown(sheet, validations).formula1 = stat_formula(table)
            target = output / source.name
            workbook.save(target)
        finally:
            workbook.close()
        reviewed = load_workbook(target)
        try:
            after = cell_projection(reviewed.active)
            if any(after.get(location) != value for location, value in before.items()):
                raise ValueError(f"{table} changed an existing cell")
            if reviewed.active.data_validations != validations:
                raise ValueError(f"{table} changed unowned data validation metadata")
            if expected:
                check_rows(reviewed.active, expected, append=False)
        finally:
            reviewed.close()
        if hashlib.sha256(source.read_bytes()).hexdigest() != source_digest:
            raise ValueError(f"{table} source changed during clean-target authoring")
        print(f"{target}: preserved {len(before)} cells; source SHA-256 {source_digest}")


def main() -> None:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--write-clean", type=Path)
    mode.add_argument("--check", action="store_true")
    arguments = parser.parse_args()
    if arguments.write_clean:
        write_clean(arguments.write_clean)
    else:
        for table, expected in records().items():
            workbook = load_workbook(DATA / f"{table}.xlsx")
            try:
                check_rows(workbook.active, expected, append=False)
            finally:
                workbook.close()
        for table in STAT_TABLES:
            workbook = load_workbook(DATA / f"{table}.xlsx")
            try:
                if stat_dropdown(workbook.active).formula1 != stat_formula(table):
                    raise ValueError(f"{table} stat dropdown differs from Sora template")
            finally:
                workbook.close()
        print("Dedicated Elation authoring probes match; no content coverage implied.")


if __name__ == "__main__":
    main()
