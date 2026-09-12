#!/usr/bin/env python3
"""Render Divergent Universe workbook review ranges through HTML and Chromium."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook


WORKBOOKS = (
    "DivergentUniverse.xlsx",
    "DivergentUniverseBindings.xlsx",
    "DivergentUniverseReview.xlsx",
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def color(value: object, fallback: str) -> str:
    rgb = getattr(value, "rgb", None)
    if isinstance(rgb, str) and len(rgb) >= 6:
        return f"#{rgb[-6:]}"
    return fallback


def render_cell(cell: object) -> str:
    value = getattr(cell, "value", None)
    text = "" if value is None else str(value)
    fill = color(getattr(getattr(cell, "fill", None), "fgColor", None), "#ffffff")
    font = getattr(cell, "font", None)
    foreground = color(getattr(font, "color", None), "#111827")
    weight = "700" if getattr(font, "bold", False) else "400"
    return (
        f'<td style="background:{fill};color:{foreground};font-weight:{weight}">'
        f'<div>{html.escape(text)}</div></td>'
    )


def sheet_html(sheet: object, field_count: int, title: str) -> str:
    widths = []
    for column in range(1, field_count + 1):
        letter = sheet.cell(row=1, column=column).column_letter
        authored = sheet.column_dimensions[letter].width or 10
        widths.append(max(90, min(480, int(authored * 8))))
    rows = []
    for row in sheet.iter_rows(min_row=1, max_row=12, max_col=field_count):
        rows.append("<tr>" + "".join(render_cell(cell) for cell in row) + "</tr>")
    columns = "".join(f'<col style="width:{width}px">' for width in widths)
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
html,body{{margin:0;background:#e5e7eb;font-family:Segoe UI,Microsoft YaHei,sans-serif}}
main{{padding:16px}} h1{{font-size:18px;margin:0 0 10px;color:#111827}}
.frame{{display:inline-block;background:white;padding:8px;box-shadow:0 1px 5px #64748b}}
table{{border-collapse:collapse;table-layout:fixed;font-size:12px}}
td{{border:1px solid #94a3b8;padding:4px;vertical-align:top;overflow:hidden}}
td div{{white-space:pre-wrap;overflow:hidden;max-height:96px;line-height:16px}}
</style></head><body><main><h1>{html.escape(title)}</h1>
<div class="frame"><table><colgroup>{columns}</colgroup>{''.join(rows)}</table></div>
</main></body></html>"""


def screenshot(browser: Path, page: Path, target: Path, size: str, profile: Path) -> None:
    result = subprocess.run(
        [
            str(browser),
            "--headless=new",
            "--disable-gpu",
            "--hide-scrollbars",
            "--no-first-run",
            "--allow-file-access-from-files",
            f"--user-data-dir={profile}",
            f"--window-size={size}",
            f"--screenshot={target}",
            page.resolve().as_uri(),
        ],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if result.returncode != 0 or not target.exists():
        raise RuntimeError(
            f"browser screenshot failed for {page}:\n{result.stdout}\n{result.stderr}"
        )


def contact_html(images: list[Path], title: str) -> str:
    cards = "".join(
        f'<figure><img src="{image.resolve().as_uri()}"><figcaption>'
        f"{html.escape(image.stem)}</figcaption></figure>"
        for image in images
    )
    return f"""<!doctype html><html><head><meta charset="utf-8">
<style>
html,body{{margin:0;background:#111827;color:white;font-family:Segoe UI,sans-serif}}
h1{{height:34px;margin:8px 16px;font-size:20px}}
main{{display:grid;grid-template-columns:repeat(4,580px);grid-auto-rows:745px;gap:8px;padding:8px}}
figure{{margin:0;background:#f8fafc;border:1px solid #64748b;display:flex;flex-direction:column}}
img{{width:578px;height:710px;object-fit:contain;background:white}}
figcaption{{height:25px;color:#111827;font-size:11px;overflow:hidden;padding:3px}}
</style></head><body><h1>{html.escape(title)}</h1><main>{cards}</main></body></html>"""


def browser_version(browser: Path) -> str:
    if os.name == "nt":
        escaped = str(browser).replace("'", "''")
        command = (
            f"(Get-Item -LiteralPath '{escaped}').VersionInfo.ProductVersion"
        )
        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", command],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        return f"Microsoft Edge {result.stdout.strip()}"
    result = subprocess.run(
        [str(browser), "--version"],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return result.stdout.strip()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", nargs="?", default=".")
    parser.add_argument("output")
    parser.add_argument("--browser", required=True)
    args = parser.parse_args()
    root = Path(args.root).resolve()
    output = Path(args.output).resolve()
    browser = Path(args.browser).resolve()
    if output.exists():
        raise FileExistsError(f"refusing to overwrite {output}")
    if not browser.is_file():
        raise FileNotFoundError(browser)
    output.mkdir(parents=True)
    schema = json.loads(
        (root / "config/divergent-universe-generated/schema.lock").read_text(
            encoding="utf-8"
        )
    )["schema"]
    rendered = []
    ordinal = 0
    with tempfile.TemporaryDirectory(prefix="starclock-du-browser-") as temporary:
        profile = Path(temporary) / "profile"
        for workbook_name in WORKBOOKS:
            workbook = load_workbook(
                root / "config/divergent-universe/data" / workbook_name,
                read_only=False,
                data_only=False,
            )
            tables = [
                table
                for table in schema["tables"]
                if table["source"]["file"] == workbook_name
            ]
            for table in tables:
                ordinal += 1
                sheet_name = table["source"]["sheet"]
                stem = (
                    f"{ordinal:02d}-{Path(workbook_name).stem}-{sheet_name}"
                )
                page = output / f"{stem}.html"
                image = output / f"{stem}.png"
                page.write_text(
                    sheet_html(
                        workbook[sheet_name],
                        len(table["fields"]),
                        f"{workbook_name} / {sheet_name} / A1:{len(table['fields'])}x12",
                    ),
                    encoding="utf-8",
                )
                screenshot(browser, page, image, "6000,1200", profile)
                rendered.append(
                    {
                        "file": workbook_name,
                        "sheet": sheet_name,
                        "range": f"A1:{len(table['fields'])}x12",
                        "image": image.name,
                        "sha256": sha256(image),
                    }
                )
            workbook.close()
        contacts = []
        for index in range((len(rendered) + 7) // 8):
            images = [
                output / item["image"]
                for item in rendered[index * 8 : (index + 1) * 8]
            ]
            page = output / f"contact-{index + 1:02d}.html"
            image = output / f"contact-{index + 1:02d}.png"
            page.write_text(
                contact_html(images, f"Divergent Universe sheets {index * 8 + 1}-{min(index * 8 + 8, len(rendered))}"),
                encoding="utf-8",
            )
            screenshot(browser, page, image, "2400,1600", profile)
            contacts.append({"image": image.name, "sha256": sha256(image)})
    manifest = {
        "schema_revision": "starclock.divergent-universe-visual-render.v1",
        "renderer": {
            "name": "openpyxl-html-chromium",
            "version": browser_version(browser),
            "range_policy": "rows 1-12 across every used schema column",
        },
        "sheet_count": len(rendered),
        "sheets": rendered,
        "contact_sheets": contacts,
    }
    manifest_path = output / "render-manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"Rendered {len(rendered)} sheets and {len(contacts)} contact sheets; "
        f"manifest_sha256={sha256(manifest_path)}"
    )


if __name__ == "__main__":
    main()
